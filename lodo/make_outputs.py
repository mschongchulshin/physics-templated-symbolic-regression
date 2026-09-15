import os, sys
import numpy as np, pandas as pd
import matplotlib.pyplot as plt
from matplotlib.lines import Line2D
from matplotlib.colors import TwoSlopeNorm
from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill, Alignment

HERE = os.path.dirname(os.path.abspath(__file__))
import ncstyle as N
N.use_nc_style()
FIG = f"{HERE}/figures"; os.makedirs(FIG, exist_ok=True)

P = pd.read_csv(f"{HERE}/results_ptsr.csv")
B = pd.read_csv(f"{HERE}/results_baselines.csv")
M = pd.read_pickle(f"{HERE}/verified_M.pkl")
V = pd.read_pickle(f"{HERE}/verified_V.pkl")

ORDER = [
    ("T1|AlCrFeMnNi|CAST|C|Zhang 2019", "F1", "Tier 1"),
    ("T2|CoCrFeMoNi|WROUGHT|T|Ming 2019 + Ming 2017", "F2", "Tier 2"),
    ("T2|CoCrFeMoNi|WROUGHT|T|Wei 2018 + Bae 2019", "F3", "Tier 2"),
    ("T3|AlCoCrFeMnNi|CAST|T|He 2014", "F4", "Tier 3"),
    ("T3|AlCoCrFeMoNi|CAST|C|Wang 2008 + Zhu 2011 + Zhu 2010 + Tian 2019", "F5", "Tier 3"),
    ("T3|HfMoNbTaTiZr|CAST|C|Juan 2016", "F6", "Tier 3"),
    ("T0|CoCrFeNi|WROUGHT|T|Wei 2018", "F7", "Extra"),
]
TT = {"C": "compression", "T": "tension"}
fold_info = []
for fid, short, tier in ORDER:
    _, sysname, proc, tt, lab = fid.split("|")
    mm = M[M.fold_id == fid]
    te = V.loc[mm[mm.role == "test"].row]
    tr = V.loc[mm[mm.role == "train"].row]
    test_papers = "; ".join(f"{p} ({n})" for p, n in te.lab_name.value_counts().sort_index().items())
    train_papers = "; ".join(f"{p} ({n})" for p, n in tr.lab_name.value_counts().sort_index().items())
    fold_info.append(dict(fold_id=fid, Fold=short, Tier=tier, Alloy_system=sysname, Processing=proc,
                          Test_type=TT[tt], n_train=len(tr), n_test=len(te), Test_papers=test_papers,
                          Train_papers=train_papers,
                          label=f"{short} {sysname} {proc} {TT[tt][0].upper()} · {lab.split(' + ')[0]}{' +' if ' + ' in lab else ''} ({len(tr)}/{len(te)})"))
FI = pd.DataFrame(fold_info)
SH = dict(zip(FI.fold_id, FI.Fold))
LBL = dict(zip(FI.Fold, FI.label))
LEG = {r.Fold: f"{r.Fold} {r.Alloy_system} {r.Processing} ({r.Test_type})" for r in FI.itertuples()}
NAIVE = {}
for fid, short, _ in ORDER:
    mm = M[M.fold_id == fid]
    ytr = V.loc[mm[mm.role == "train"].row].YS_MPa.values
    yte = V.loc[mm[mm.role == "test"].row].YS_MPa.values
    NAIVE[short] = dict(mae=float(np.abs(yte - ytr.mean()).mean()), test_sd=float(yte.std()),
                        test_range=f"{yte.min():.0f}-{yte.max():.0f}", train_range=f"{ytr.min():.0f}-{ytr.max():.0f}")
FI["Test_YS_range_MPa"] = FI.Fold.map(lambda f: NAIVE[f]["test_range"])
FI["Test_YS_SD_MPa"] = FI.Fold.map(lambda f: round(NAIVE[f]["test_sd"], 1))
FI["Train_YS_range_MPa"] = FI.Fold.map(lambda f: NAIVE[f]["train_range"])
FI["Naive_train_mean_MAE_MPa"] = FI.Fold.map(lambda f: round(NAIVE[f]["mae"], 1))

agg = (P.groupby(["fold_id", "model"])
         .agg(train_mean=("train_r2", "mean"), test_mean=("test_r2", "mean"),
              test_median=("test_r2", "median"), test_min=("test_r2", "min"),
              test_max=("test_r2", "max"), mae_mean=("test_mae", "mean"),
              n_ok=("test_r2", lambda s: int(s.notna().sum())))
         .reset_index())
fair, oracle = [], []
for fid, g in agg.groupby("fold_id"):
    gv = g.dropna(subset=["train_mean"])
    if len(gv):
        f = gv.sort_values(["train_mean", "model"], ascending=[False, True]).iloc[0]
        fair.append(dict(fold_id=fid, template=f.model, test_r2=f.test_mean, train_r2=f.train_mean,
                         lo=f.test_min, hi=f.test_max, mae=f.mae_mean))
    go = g.dropna(subset=["test_mean"])
    if len(go):
        o = go.sort_values(["test_mean", "model"], ascending=[False, True]).iloc[0]
        oracle.append(dict(fold_id=fid, template=o.model, test_r2=o.test_mean, lo=o.test_min, hi=o.test_max, mae=o.mae_mean))
FAIR = pd.DataFrame(fair); ORA = pd.DataFrame(oracle)

BA = (B.groupby(["fold_id", "model"])
        .agg(test_r2=("test_r2", "median"), lo=("test_r2", "min"), hi=("test_r2", "max"),
             mae=("test_mae", "median"), train_r2=("train_r2", "median"))
        .reset_index())

NAME = {"PT-SR_fair": "PT-SR (this work)", "PT-SR_oracle": "PT-SR capacity ceiling",
        "PySR_no_template": "PySR (Cranmer 2023)", "SISSO": "SISSO (Ouyang 2018)",
        "Wu_gplearn_RFR": "Wu gplearn-RFR (2024)", "Liu_MLP": "Liu MLP (2024)", "Liu_RF": "Liu RF (2024)",
        "LESets_GNN": "LESets GNN (Zhang 2024)", "Jain_DNN": "Jain DNN (2026)",
        "SciRep25_Transformer": "Korkmaz Transformer (2025)", "JMI_Stacking": "JMI Stacking (2024)"}
long = [FAIR.assign(model="PT-SR_fair"), ORA.assign(model="PT-SR_oracle"), BA]
L = pd.concat(long, ignore_index=True)
L["Fold"] = L.fold_id.map(SH)
L = L[L.Fold.notna()]

R2 = L.pivot_table(index="model", columns="Fold", values="test_r2").reindex(columns=FI.Fold)
MAE = L.pivot_table(index="model", columns="Fold", values="mae").reindex(columns=FI.Fold)
LO = L.pivot_table(index="model", columns="Fold", values="lo").reindex(columns=FI.Fold)
HI = L.pivot_table(index="model", columns="Fold", values="hi").reindex(columns=FI.Fold)
THR = 0.9
naive_vec = pd.Series({f: NAIVE[f]["mae"] for f in FI.Fold})
SKILL = 1.0 - MAE.div(naive_vec, axis=1)
summ = pd.DataFrame({
    "Method": [NAME.get(m, m) for m in R2.index],
    "Folds_R2_above_0.9": (R2 > THR).sum(axis=1).values,
    "Folds_R2_above_0": (R2 > 0).sum(axis=1).values,
    "Folds_MAE_better_than_naive": (SKILL > 0).sum(axis=1).values,
    "Median_MAE_skill_vs_naive": SKILL.median(axis=1).values,
    "Median_test_R2": R2.median(axis=1).values,
    "Median_test_MAE_MPa": MAE.median(axis=1).values,
    "Widest_seed_spread_R2": (HI - LO).max(axis=1).values,
}, index=R2.index).sort_values(["Folds_MAE_better_than_naive", "Median_MAE_skill_vs_naive"], ascending=False)

table = R2.copy(); table.index = [NAME.get(m, m) for m in table.index]
table.columns = [LBL[c] for c in table.columns]
table = table.loc[summ.Method]
mae_t = MAE.copy(); mae_t.index = [NAME.get(m, m) for m in mae_t.index]; mae_t.columns = table.columns; mae_t = mae_t.loc[summ.Method]
skill_t = SKILL.copy(); skill_t.index = [NAME.get(m, m) for m in skill_t.index]; skill_t.columns = table.columns; skill_t = skill_t.loc[summ.Method]
spread = (HI - LO); spread.index = [NAME.get(m, m) for m in spread.index]; spread.columns = table.columns; spread = spread.loc[summ.Method]
tmpl = FAIR.assign(Fold=FAIR.fold_id.map(SH))[["Fold", "template", "train_r2", "test_r2", "lo", "hi", "mae"]]
tmpl = tmpl.merge(ORA.assign(Fold=ORA.fold_id.map(SH))[["Fold", "template", "test_r2"]].rename(columns={"template": "oracle_template", "test_r2": "oracle_test_r2"}), on="Fold")
tmpl.columns = ["Fold", "Fair_template (max mean train R2)", "Fair_train_R2", "Fair_test_R2 (mean of 5 seeds)", "Seed_min", "Seed_max", "Test_MAE_MPa", "Oracle_template", "Oracle_test_R2"]
per_tmpl = agg.assign(Fold=agg.fold_id.map(SH)).drop(columns="fold_id")
runs = pd.concat([P.assign(family="PT-SR template"), B.assign(family="baseline")], ignore_index=True)
runs["Fold"] = runs.fold_id.map(SH)
rows = M.merge(V[["doi", "lab_name", "formula", "processing", "test_type", "T_C", "YS_MPa", "verdict"]], left_on="row", right_index=True)
rows["Fold"] = rows.fold_id.map(SH); rows = rows[rows.Fold.notna()].sort_values(["Fold", "role", "lab_name"])
rows = rows[["Fold", "role", "lab_name", "doi", "formula", "processing", "test_type", "T_C", "YS_MPa", "verdict"]]
readme = pd.DataFrame({"Item": ["Data", "Folds", "Models", "Features", "PT-SR selection", "Baseline aggregation", "Caution"],
    "Description": [
        "LODO_verified_dataset.xlsx: 138 yield-strength values verified against the original articles (DOI per row).",
        "7 leave-one-laboratory-out folds within one alloy system, one processing route and one test type. F1 passes every rule; F2-F3 extrapolate in composition; F4-F6 have no full-system composition in training; F7 (Wei 2018 CoCrFeNi) trains on a single composition.",
        "PT-SR 9 templates (repo lodo/ptsr_templates.py) and 9 baselines (repo baselines/lodo_comparison/run.py), code and hyperparameters unchanged; seeds 0-4, PySR niterations 30.",
        "Element fractions of the alloy system (zero for absent elements), S_mix, H_mix, delta, VEC, dChi, r_avg, chi_avg, T (K), as in repo lodo/features.py.",
        "Fair: template with the highest mean training R2 over 5 seeds, reporting its mean test R2. Oracle: highest mean test R2 (upper bound, uses test data).",
        "Median test R2 over seeds (Liu RF and SISSO single seed), as in repo lodo/run_baselines.py.",
        "Test sets hold 3-11 measurements and several span less than 110 MPa (F1, F5, F6), so R2 becomes strongly negative for errors of a few tens of MPa. MAE skill = 1 - MAE(model)/MAE(predicting the training mean) is reported alongside: > 0 beats the naive reference, 1 is perfect."]})
out = f"{HERE}/LODO_benchmark.xlsx"
with pd.ExcelWriter(out, engine="openpyxl") as w:
    readme.to_excel(w, sheet_name="README", index=False)
    summ.to_excel(w, sheet_name="summary_by_method", index=False)
    table.to_excel(w, sheet_name="SuppTable_test_R2")
    mae_t.to_excel(w, sheet_name="SuppTable_test_MAE")
    skill_t.to_excel(w, sheet_name="SuppTable_MAE_skill")
    spread.to_excel(w, sheet_name="seed_spread_R2")
    FI.drop(columns=["fold_id", "label"]).to_excel(w, sheet_name="folds", index=False)
    tmpl.to_excel(w, sheet_name="PTSR_selection", index=False)
    per_tmpl.to_excel(w, sheet_name="PTSR_per_template", index=False)
    runs.to_excel(w, sheet_name="all_runs", index=False)
    rows.to_excel(w, sheet_name="fold_rows", index=False)
wb = load_workbook(out)
for ws in wb.worksheets:
    for c in ws[1]:
        c.font = Font(bold=True); c.fill = PatternFill("solid", fgColor="DDEBF7"); c.alignment = Alignment(wrap_text=True, vertical="top")
    ws.freeze_panes = "B2"
    for col in ws.columns:
        ws.column_dimensions[col[0].column_letter].width = min(48, max(9, *(len(str(c.value or "")) for c in col[:40])) + 1)
    for row in ws.iter_rows(min_row=2):
        for c in row:
            if isinstance(c.value, float): c.number_format = "0.000"
wb["README"].column_dimensions["B"].width = 140
wb.save(out)

KIND = {"PT-SR_fair": "ptsr", "PySR_no_template": "sr", "SISSO": "sr"}
methods = [m for m in summ.index if m != "PT-SR_oracle"][::-1]
col = lambda m: {"ptsr": N.C["ptsr"], "sr": N.C["free"]}.get(KIND.get(m), N.C["sota"])
nF = len(FI); mk = ["o", "s", "D", "^", "v", "P", "X"]
fig, axes = plt.subplots(1, 3, figsize=(N.W2 * 0.99, 2.9), gridspec_kw=dict(width_ratios=[0.62, 1.0, 1.0], wspace=0.08))
axC, axK, axR = axes
y = np.arange(len(methods))
for i, m in enumerate(methods):
    lead = m == "PT-SR_fair"; n = int(summ.loc[m, "Folds_MAE_better_than_naive"])
    axC.barh(i, n, height=0.64, color=col(m), alpha=1 if lead else 0.72, lw=0, zorder=2)
    axC.text(n + 0.12, i, f"{n}/{nF}", va="center", fontsize=7.2, color=col(m), fontweight="bold" if lead else "normal")
axC.set_yticks(y); axC.set_yticklabels([NAME[m].split(" (")[0] for m in methods])
for t, m in zip(axC.get_yticklabels(), methods):
    if m == "PT-SR_fair": t.set_fontweight("bold"); t.set_color(N.C["ptsr"])
axC.set_xlim(0, nF + 1.4); axC.set_xticks(range(nF + 1)); axC.set_xlabel("Folds beating\ntraining mean (MAE)")
def dots(ax, T, floor, lo_lab):
    for i, m in enumerate(methods):
        for j, f in enumerate(FI.Fold):
            v = T.loc[m, f]
            if pd.isna(v): continue
            ax.scatter(max(v, floor), i + (j - 3) * 0.08, s=15, marker=mk[j], color=col(m), alpha=0.9, lw=0, zorder=3, clip_on=False)
    ax.set_yticks(y); ax.set_yticklabels([])
axK.axvspan(0, 1.05, color="#eef3f9", zorder=0); axK.axvline(0, color="#a9b3bd", lw=0.7, ls=(0, (3, 2.5)))
dots(axK, SKILL, -3.0, "≤−3"); axK.set_xlim(-3.15, 1.08); axK.set_xticks([-3, -2, -1, 0, 1]); axK.set_xticklabels(["≤−3", "−2", "−1", "0", "1"])
axK.set_xlabel("MAE skill vs training mean")
axR.axvspan(THR, 1.05, color="#eef3f9", zorder=0); axR.axvline(THR, color="#a9b3bd", lw=0.7, ls=(0, (3, 2.5))); axR.axvline(0, color="#c6cbd1", lw=0.6)
dots(axR, R2, -2.0, "≤−2"); axR.set_xlim(-2.15, 1.08); axR.set_xticks([-2, -1, 0, THR, 1]); axR.set_xticklabels(["≤−2", "−1", "0", "0.9", "1"])
axR.set_xlabel("Test R²")
for a_ in axes:
    a_.set_ylim(-0.7, len(methods) - 0.3); a_.xaxis.grid(True, color=N.C["grid"], lw=0.5); a_.set_axisbelow(True)
h = [Line2D([], [], marker=mk[j], ls="none", ms=4.5, color="#555555", label=LEG[f]) for j, f in enumerate(FI.Fold)]
axK.legend(handles=h, loc="upper center", bbox_to_anchor=(0.55, -0.24), ncol=4, fontsize=6.3, handletextpad=0.2, columnspacing=0.9)
N.panel_label(axC, "a", dx=-0.62, dy=1.07); N.panel_label(axK, "b", dx=-0.02, dy=1.07); N.panel_label(axR, "c", dx=-0.02, dy=1.07)
fig.savefig(f"{FIG}/Fig5_lodo.svg", bbox_inches="tight"); fig.savefig(f"{FIG}/Fig5_lodo.png", bbox_inches="tight", dpi=300)
plt.close(fig)

rows_order = list(summ.index)
fig, axs = plt.subplots(2, 1, figsize=(N.W2, 6.4), gridspec_kw=dict(hspace=0.12))
for ax, T, title, vmin, fmt in [(axs[0], R2, "Test R² (colour clipped to [−1, 1])", -1, "{:.2f}"), (axs[1], SKILL, "MAE skill vs training-mean predictor (colour clipped to [−1, 1])", -1, "{:.2f}")]:
    H = T.loc[rows_order]
    im = ax.imshow(H.clip(-1, 1).values, cmap="RdBu", norm=TwoSlopeNorm(vcenter=0.0, vmin=-1, vmax=1), aspect="auto")
    for i in range(H.shape[0]):
        for j in range(H.shape[1]):
            v = H.values[i, j]
            s = "n/a" if pd.isna(v) else (fmt.format(v) if v > -10 else f"{v:.0f}")
            ax.text(j, i, s, ha="center", va="center", fontsize=6, color="white" if (not pd.isna(v) and abs(np.clip(v, -1, 1)) > 0.6) else "#222222")
    ax.set_yticks(range(H.shape[0])); ax.set_yticklabels([NAME[m] for m in H.index], fontsize=6.8)
    ax.set_xticks(range(H.shape[1])); ax.set_xticklabels([] if ax is axs[0] else [LBL[f] for f in H.columns], rotation=35, ha="right", fontsize=6.2)
    for s_ in ax.spines.values(): s_.set_visible(False)
    cb = fig.colorbar(im, ax=ax, fraction=0.022, pad=0.01); cb.set_label(title, fontsize=6.5)
N.panel_label(axs[0], "a", dx=-0.3, dy=1.04); N.panel_label(axs[1], "b", dx=-0.3, dy=1.04)
fig.savefig(f"{FIG}/SuppFig_lodo_perfold.svg", bbox_inches="tight"); fig.savefig(f"{FIG}/SuppFig_lodo_perfold.png", bbox_inches="tight", dpi=300)
plt.close(fig)

pd.set_option("display.width", 250)
print(out)
print(summ.round(3).to_string())
print(table.round(2).to_string()); print(skill_t.round(2).to_string()); print(FI[["Fold","Test_YS_range_MPa","Test_YS_SD_MPa","Train_YS_range_MPa","Naive_train_mean_MAE_MPa"]].to_string(index=False))
print(tmpl.round(3).to_string(index=False))
